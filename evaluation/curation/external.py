#!/usr/bin/env python3
"""Normalize the independent Excel question corpus into one reviewable YAML file."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, cast

import yaml
from openpyxl import load_workbook

from shared.classifier import CLASSIFIER_PROMPT_TEMPLATE

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = ROOT / "tests/evaluation/external_questions.yaml"
VALID_INTENTS = {"specific_plan", "broad_plans", "generic_info"}
VALID_MARKERS = {"CS", "C", "CE", "CSE"}

# This exact question was present in a historical prompt before prompt cleanup.
# Keep the exclusion reproducible even after that example is removed.
HISTORICAL_PROMPT_OVERLAPS = {"can i get chiropractic care with this plan"}


class CleanDumper(yaml.SafeDumper):
    """Use readable block scalars for multiline reference answers."""


def _represent_string(dumper: yaml.SafeDumper, value: str):
    style = "|" if "\n" in value else None
    return dumper.represent_scalar("tag:yaml.org,2002:str", value, style=style)


CleanDumper.add_representer(str, cast(Any, _represent_string))


def _normalize_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def _numbered_row(value: Any) -> bool:
    return isinstance(value, int | float) or str(value or "").strip().isdigit()


def _optional_text(value: Any) -> str | None:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in str(value or "").splitlines()]
    compacted: list[str] = []
    for line in lines:
        if line or (compacted and compacted[-1]):
            compacted.append(line)
    text = "\n".join(compacted).strip()
    return text or None


def _main_prompt_corpus() -> str:
    """Return only the prompts active in the current agent source tree."""
    agent_source = (ROOT / "src/agent.yaml").read_text()
    return _normalize_text(f"{CLASSIFIER_PROMPT_TEMPLATE} {agent_source}")


def _prompt_overlaps(question: str, main_prompt: str) -> list[str]:
    normalized = _normalize_text(question)
    if len(normalized.split()) < 3:
        return []
    labels = []
    if normalized in main_prompt:
        labels.append("main")
    if normalized in HISTORICAL_PROMPT_OVERLAPS:
        labels.append("refactor_pre_cleanup")
    return labels


def _parse_context(raw: Any) -> dict[str, Any]:
    try:
        source = json.loads(raw) if isinstance(raw, str) else raw
    except json.JSONDecodeError as exc:
        raise ValueError("context is not valid JSON") from exc
    if not isinstance(source, dict):
        raise ValueError("context must be an object")

    application = source.get("application_context", {})
    prospect = source.get("prospect_context", {})
    user = source.get("user_context", {})
    available = application.get("available_plans", [])
    recommended = application.get("recommended_plans", [])
    if not isinstance(available, list) or not available:
        raise ValueError("context must contain available plans")
    if not isinstance(recommended, list):
        raise ValueError("recommended plans must be a list")

    return {
        "user_brand": user.get("brand", "ABC"),
        "user_zip_code": user.get("zip_code", ""),
        "user_county_code": user.get("county_code", ""),
        "user_state_code": user.get("state_code", ""),
        "user_requested_eff_date": user.get("requested_eff_date", ""),
        "user_dsnp_eligibility": user.get("dsnp_eligibility", ""),
        "user_applicant_count": user.get("applicant_count", 1),
        "user_current_plan": user.get("current_plan", ""),
        "user_language": user.get("language", "en"),
        "application_exchange_indicator": application.get("exchange_indicator", ""),
        "application_market_segment": application.get("market_segment", ""),
        "application_available_plans": available,
        "application_recommended_plans": recommended,
        "application_current_page": application.get("current_page", ""),
        "prospect_type": prospect.get("prospect_type", "prospect"),
    }


def _context_key(context: dict[str, Any]) -> str:
    encoded = json.dumps(context, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    return hashlib.sha256(encoded.encode()).hexdigest()


def _context_base_name(context: dict[str, Any]) -> str:
    segment = str(context["application_market_segment"]).casefold()
    segment = "individual" if segment == "ind" else segment
    prospect = str(context["prospect_type"]).casefold() or "unknown"
    suffix = "_current_plan" if context.get("user_current_plan") else ""
    return f"{segment}_{prospect}{suffix}"


def _context_registry(rows: list[dict[str, Any]]) -> tuple[dict[str, dict], dict[str, str]]:
    unique: dict[str, dict[str, Any]] = {}
    for row in rows:
        unique.setdefault(_context_key(row["context"]), row["context"])

    contexts: dict[str, dict] = {}
    names_by_hash: dict[str, str] = {}
    counts: Counter[str] = Counter()
    for digest, context in unique.items():
        base = _context_base_name(context)
        counts[base] += 1
        name = base if counts[base] == 1 else f"{base}_{counts[base]:02d}"
        contexts[name] = context
        names_by_hash[digest] = name
    return contexts, names_by_hash


def _read_rows(source: Path, main_prompt: str) -> list[dict[str, Any]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    try:
        for sheet in workbook:
            if sheet.title not in {"Individual", "Medicare"}:
                continue
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=2, max_col=7, values_only=True), start=2
            ):
                number, nature, question, answer, context, rag, marker = values
                if not question or not _numbered_row(number):
                    continue
                marker = str(marker or "CSE").strip().upper()
                if marker not in VALID_MARKERS:
                    marker = "CSE"
                source_intent = str(nature or "").strip() or None
                if source_intent not in VALID_INTENTS:
                    source_intent = None
                question = re.sub(r"\s+", " ", str(question)).strip()
                overlaps = _prompt_overlaps(question, main_prompt)
                reference_answer = _optional_text(answer)
                status = (
                    "excluded_prompt_overlap"
                    if overlaps
                    else (
                        "needs_review"
                        if reference_answer is None or source_intent is None
                        else "eligible"
                    )
                )
                rows.append(
                    {
                        "sheet": sheet.title,
                        "source_row": row_number,
                        "source_question_number": int(float(str(number))),
                        "question": question,
                        "reference_answer": reference_answer,
                        "source_intent": source_intent,
                        "source_rag_expected": str(rag or "").strip().casefold() == "yes",
                        "marker": marker,
                        "context": _parse_context(context),
                        "prompt_overlap": overlaps,
                        "evaluation_status": status,
                    }
                )
    finally:
        workbook.close()
    return rows


def _group_rows(rows: list[dict[str, Any]]) -> list[tuple[str, list[dict[str, Any]]]]:
    groups: list[tuple[str, list[dict[str, Any]]]] = []
    for sheet in ("Individual", "Medicare"):
        active: list[dict[str, Any]] = []
        for row in (item for item in rows if item["sheet"] == sheet):
            marker = row["marker"]
            if marker == "CS":
                if active:
                    groups.append((sheet, active))
                active = [row]
            elif marker == "C":
                if active:
                    active.append(row)
                else:
                    groups.append((sheet, [row]))
            elif marker == "CE":
                if active:
                    active.append(row)
                    groups.append((sheet, active))
                    active = []
                else:
                    groups.append((sheet, [row]))
            else:
                if active:
                    groups.append((sheet, active))
                    active = []
                groups.append((sheet, [row]))
        if active:
            groups.append((sheet, active))
    return groups


def _conversation_status(turns: list[dict[str, Any]]) -> str:
    statuses = {turn["evaluation_status"] for turn in turns}
    if "excluded_prompt_overlap" in statuses:
        return "excluded_prompt_overlap"
    if "needs_review" in statuses:
        return "needs_review"
    return "eligible"


def _build_conversations(
    groups: list[tuple[str, list[dict[str, Any]]]], names_by_hash: dict[str, str]
) -> list[dict[str, Any]]:
    conversations = []
    sheet_counts: Counter[str] = Counter()
    for sheet, source_turns in groups:
        sheet_key = sheet.casefold()
        sheet_counts[sheet_key] += 1
        conversation_id = f"external_{sheet_key}_{sheet_counts[sheet_key]:03d}"
        default_context = names_by_hash[_context_key(source_turns[0]["context"])]
        turns = []
        for index, source in enumerate(source_turns, start=1):
            context_name = names_by_hash[_context_key(source["context"])]
            turn = {
                "id": f"{conversation_id}_turn_{index:02d}",
                "source_row": source["source_row"],
                "source_question_number": source["source_question_number"],
                "question": source["question"],
                "reference_answer": source["reference_answer"],
                "source_intent": source["source_intent"],
                "source_rag_expected": source["source_rag_expected"],
                "evaluation_status": source["evaluation_status"],
                "prompt_overlap": source["prompt_overlap"],
            }
            if context_name != default_context:
                turn["context"] = context_name
            turns.append(turn)
        conversations.append(
            {
                "id": conversation_id,
                "market": sheet,
                "context": default_context,
                "evaluation_status": _conversation_status(turns),
                "turns": turns,
            }
        )
    return conversations


def _statistics(
    rows: list[dict[str, Any]], contexts: dict[str, dict], conversations: list[dict[str, Any]]
) -> dict[str, Any]:
    turn_statuses = Counter(row["evaluation_status"] for row in rows)
    conversation_statuses = Counter(item["evaluation_status"] for item in conversations)
    return {
        "turns": len(rows),
        "unique_questions": len({_normalize_text(row["question"]) for row in rows}),
        "conversations": len(conversations),
        "contexts": len(contexts),
        "turns_by_market": dict(Counter(row["sheet"] for row in rows)),
        "turns_by_status": dict(turn_statuses),
        "conversations_by_status": dict(conversation_statuses),
        "missing_reference_answers": sum(row["reference_answer"] is None for row in rows),
        "missing_source_intents": sum(row["source_intent"] is None for row in rows),
    }


def curate(source: Path) -> dict[str, Any]:
    main_prompt = _main_prompt_corpus()
    rows = _read_rows(source, main_prompt)
    contexts, names_by_hash = _context_registry(rows)
    conversations = _build_conversations(_group_rows(rows), names_by_hash)
    return {
        "schema_version": 1,
        "source": {
            "file": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "sheets": ["Individual", "Medicare"],
        },
        "curation": {
            "purpose": "Independent behavioral evaluation against local watsonx Orchestrate",
            "answer_policy": (
                "reference_answer is retained for review and is never an exact-response oracle"
            ),
            "intent_policy": (
                "source_intent is an imported annotation and is not authoritative routing truth"
            ),
            "context_policy": (
                "nested backend contexts are deduplicated and normalized to local WXO variables"
            ),
            "prompt_overlap_policy": (
                "exclude entire conversations containing questions copied into either "
                "baseline prompt"
            ),
        },
        "statistics": _statistics(rows, contexts, conversations),
        "contexts": contexts,
        "conversations": conversations,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if not args.source.is_file():
        parser.error(f"source workbook not found: {args.source}")

    corpus = curate(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        yaml.dump(
            corpus,
            Dumper=CleanDumper,
            allow_unicode=True,
            default_flow_style=False,
            sort_keys=False,
            width=100,
        )
    )
    stats = corpus["statistics"]
    print(
        f"Wrote {stats['turns']} turns in {stats['conversations']} conversations "
        f"with {stats['contexts']} contexts to {args.output}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
